"""Dumps every paragraph and table cell in a .docx to one JSON line each.

Used to build and check the ground-truth annotations for the evaluation
report (eval/ground_truth.jsonl) -- having a stable id per text unit lets
annotations reference exactly one paragraph or cell, independent of how
the redactor itself walks the document.

Usage:
    python scripts/extract_text.py "input.docx" units.jsonl
"""

import json
import sys

from docx import Document


def extract_units(docx_path: str):
    """Yields {"id": ..., "text": ...} for every non-empty paragraph and table cell."""
    document = Document(docx_path)

    for paragraph_index, paragraph in enumerate(document.paragraphs):
        if paragraph.text.strip():
            yield {"id": f"p{paragraph_index}", "text": paragraph.text}

    for table_index, table in enumerate(document.tables):
        seen_cells = set()
        for row_index, row in enumerate(table.rows):
            for col_index, cell in enumerate(row.cells):
                if cell._tc in seen_cells:
                    continue  # skip merged-cell repeats, same as docx_io.py
                seen_cells.add(cell._tc)
                if cell.text.strip():
                    yield {
                        "id": f"t{table_index}_r{row_index}_c{col_index}",
                        "text": cell.text,
                    }


if __name__ == "__main__":
    input_path, output_path = sys.argv[1], sys.argv[2]
    with open(output_path, "w", encoding="utf-8") as f:
        for unit in extract_units(input_path):
            f.write(json.dumps(unit, ensure_ascii=False) + "\n")
